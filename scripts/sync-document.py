"""Apply the approved AR scope to the canonical Grupo 18 Word document.

Run with the bundled workspace Python. This edits paragraphs in place to keep
the original covers, styles, figures, table layouts, bookmarks and TOC links.
"""

from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "docs" / "Plataforma_Vestidor3D_Grupo18.docx"

replacements = {
    "Plataforma de Comercio Electrónico Multicanal con\nVestidor 3D y Avatar Personalizado":
        "Plataforma de Comercio Electrónico Multicanal con\nVestidor de Realidad Aumentada",
    "La plataforma permitirá comprar prendas desde la web, una aplicación Android y puntos de venta físicos":
        "La plataforma permite consultar y comprar prendas desde la web, y comparte catálogo e inventario con sucursales y la aplicación Android. La experiencia prioritaria de prueba virtual usa la cámara frontal del teléfono: detecta puntos corporales en el dispositivo y superpone una imagen de prenda publicada. El visor con avatar 3D generado desde fotografías se conserva como componente legado durante la transición y se identifica como tal en sus casos y diagramas. Ambas visualizaciones son aproximadas y no garantizan talla ni ajuste físico.",
    "Desarrollar una plataforma web y móvil de comercio electrónico de prendas que integre un vestidor 3D":
        "Desarrollar una plataforma web y móvil de comercio electrónico de prendas con vestidor de realidad aumentada como experiencia prioritaria, catálogo e inventario multicanal consistentes y analítica comercial. Mantener trazabilidad entre requisitos, diseño, implementación y pruebas; conservar el avatar 3D como funcionalidad transitoria hasta validar la experiencia AR en dispositivos reales.",
    "Generar un avatar aproximado desde tres fotografías y altura":
        "Implementar la prueba virtual por cámara Android con detección corporal local, imágenes transparentes publicadas y aviso visible de aproximación; validar seguimiento, oclusión y rendimiento en dispositivos reales.",
    "Preparar prendas 3D compatibles con el cuerpo paramétrico":
        "Gestionar imágenes autorizadas por variante para realidad aumentada y conservar recursos GLB compatibles con el visor 3D legado mientras se completa la transición.",
    "La línea base comprende 25 casos de uso, 44 requisitos funcionales y ocho paquetes.":
        "La línea base original comprende 25 casos de uso, 44 requisitos funcionales y ocho paquetes. El cambio de alcance del 20 de septiembre incorpora RA01–RA08 para el vestidor por cámara, catálogo común y aplicación Android. Los casos y diagramas anteriores de avatar 3D describen el componente legado; el capítulo 10 documenta la experiencia AR prioritaria y el estado real de cada caso. Un requisito planificado no se declara implementado hasta contar con flujo, autorización, interfaz y prueba de aceptación.",
    "La primera versión mostrará una prenda a la vez sobre un cuerpo paramétrico":
        "El vestidor prioritario superpone una prenda a la vez sobre la imagen de cámara Android mediante puntos corporales calculados en el teléfono. La superposición actual usa capturas periódicas, no seguimiento continuo ni simulación física. El avatar 3D legado estima proporciones desde tres fotos y altura declarada; todavía requiere validación con personas reales y no acredita precisión antropométrica. Ninguna modalidad recomienda una talla garantizada ni sustituye la tabla de medidas.",
    "La aplicación móvil comprometida en esta línea base es Android.":
        "La aplicación móvil actual se desarrolla para Android con React Native, Expo y MediaPipe; iOS necesita implementación nativa y pruebas posteriores. La cámara AR no envía video al servidor. La web permite compra, administración y visualización del avatar 3D legado según permisos. El alcance y los límites verificables de cada modalidad se detallan en el capítulo 10.",
    "Este documento contiene especificación y diseño para desarrollo. No constituye evidencia":
        "Este documento combina la especificación original con el avance comprobado de la rama codex/avance-ar-sucursales al 21 de septiembre de 2026. Las pruebas ejecutadas y los pendientes se distinguen en el capítulo 10 y en la matriz de trazabilidad. Los capítulos anteriores conservan el diseño del avatar 3D como antecedente; no prueban por sí mismos que sus flujos estén terminados. Toda función nueva exige actualización conjunta de requisitos, casos, datos, diagramas y pruebas.",
    "Se solicitarán fotografías de frente, perfil y espalda":
        "En el componente 3D legado se solicitan fotografías de frente, perfil y espalda con ropa ajustada, postura estable y cuerpo completo visible. La interfaz recoge altura y consentimiento. El worker admite JPEG o PNG hasta 10 MB por foto; la validación con fotografías consentidas de personas y el diagnóstico de orientación siguen pendientes.",
    "El cliente revisará el avatar pendiente, confirmará o corregirá medidas":
        "En el visor legado el cliente puede revisar y aprobar un avatar, cargar una prenda GLB compatible y girar o acercar la vista. Aún faltan la corrección de medidas con regeneración y la deformación de la prenda según el cuerpo. La igualdad de plantilla evita ciertos recursos incompatibles, pero no demuestra ajuste físico.",
    "Se prioriza un pipeline paramétrico con herramientas abiertas":
        "El procesamiento del avatar legado usa herramientas abiertas: MediaPipe para puntos y máscaras, y Blender para generar el GLB. Este procesamiento no sustituye el vestidor AR prioritario. Para AR, MediaPipe ejecuta inferencia local en Android y el catálogo suministra imágenes transparentes por variante. Se medirán precisión visual, rendimiento y condiciones de iluminación con dispositivos reales; la superposición de muestra no acredita ese resultado.",
    "Esta actualización compara la línea base con la rama codex/avance-ar-sucursales al 20 de septiembre de 2026.":
        "Esta actualización compara la línea base con la rama codex/avance-ar-sucursales al 21 de septiembre de 2026. La implementación sigue incompleta: ningún caso de uso reúne todavía todas las condiciones documentadas, 18 tienen cobertura parcial y 7 permanecen pendientes. Las tablas de datos no constituyen evidencia suficiente sin lógica, autorización, interfaz y pruebas. La matriz debe revisarse tras cada cambio de código.",
    "La rama aprobó 44 pruebas de integración distribuidas en 9 suites.":
        "En la línea base de esta revisión pasaron 44 pruebas de integración distribuidas en 9 suites. Esas pruebas respaldan incrementos concretos, pero no sustituyen casos de aceptación pendientes ni validaciones con proveedores, prendas autorizadas y dispositivos reales. Nuevas correcciones se documentan con resultados de prueba separados.",
}

doc = Document(PATH)
changes = {}
for paragraph in doc.paragraphs:
    for prefix, replacement in replacements.items():
        if paragraph.text.startswith(prefix):
            if not paragraph.runs:
                raise RuntimeError(f"Paragraph without runs: {prefix}")
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""
            changes[prefix] = changes.get(prefix, 0) + 1
            break

expected = {key: (2 if key.startswith("Plataforma de Comercio Electrónico") else 1) for key in replacements}
if changes != expected:
    raise RuntimeError(f"Unexpected paragraph matches: {changes}; expected {expected}")

change_table = doc.tables[-3]
assert change_table.cell(1, 0).text == "Identificación"
for row in change_table.rows:
    if row.cells[0].text == "Función principal":
        row.cells[1].text = (
            "El diseño inicial pasó de cámara AR a avatar 3D. El cambio posterior "
            "del capítulo 10 prioriza de nuevo AR y mantiene el 3D como legado transitorio."
        )

status_table = doc.tables[-1]
assert status_table.cell(0, 0).text == "Caso"
for row in status_table.rows[1:]:
    if row.cells[0].text.startswith("CU02 "):
        row.cells[1].text = "Parcial"
        row.cells[3].text = (
            "Revalidar y revocar sesiones WebSocket; limitar intentos fallidos por cuenta "
            "y ampliar auditoría de accesos."
        )
    if row.cells[0].text.startswith("CU19 "):
        row.cells[2].text = (
            "Detector semanal corregido para comparar cada venta con cuatro semanas anteriores; "
            "evidencia y prueba de regresión."
        )
    if row.cells[0].text.startswith("CU20 "):
        row.cells[2].text = (
            "Promedio móvil por variante, ubicación y canal; horizonte posterior al corte, "
            "intervalo, persistencia y versión."
        )

doc.save(PATH)
print(f"Word actualizado: {PATH} ({sum(changes.values())} párrafos, dos tablas)")
