"""Record the deployed Lúmina catalog without disturbing the existing covers/diagrams."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / 'docs/Plataforma_Vestidor3D_Grupo18.docx'
doc = Document(PATH)
if any(p.text.startswith('11 Estado desplegado de Lúmina') for p in doc.paragraphs):
    raise SystemExit('El capítulo 11 ya existe; no se duplicó.')

anchor = '_TocLuminaCatalogo20260923'
toc_source = next(p for p in doc.paragraphs if p.text.startswith('10 Actualización de alcance'))
toc_last = next(p for p in doc.paragraphs if p.text.startswith('10.6 Trabajo necesario'))
toc = deepcopy(toc_source._p)
hyperlink = toc.find('.//' + qn('w:hyperlink'))
hyperlink.set(qn('w:anchor'), anchor)
texts = toc.findall('.//' + qn('w:t'))
texts[0].text = '11 Estado desplegado de Lúmina'
texts[-1].text = '172'
for instr in toc.findall('.//' + qn('w:instrText')):
    if instr.text and 'PAGEREF' in instr.text:
        instr.text = f' PAGEREF {anchor} \\h '
toc_last._p.addnext(toc)

break_paragraph = doc.add_paragraph()
break_paragraph.add_run().add_break(WD_BREAK.PAGE)
heading = doc.add_paragraph('11 Estado desplegado de Lúmina', style='Heading 1')
start = OxmlElement('w:bookmarkStart')
start.set(qn('w:id'), '9001')
start.set(qn('w:name'), anchor)
end = OxmlElement('w:bookmarkEnd')
end.set(qn('w:id'), '9001')
heading._p.insert(1, start)
heading._p.append(end)

doc.add_paragraph(
    'Actualización del 23 de septiembre de 2026. La instancia desplegada en Azure usa la '
    'marca Lúmina y presenta un catálogo público exclusivamente de moda femenina. '
    'Esta decisión es posterior al estado descrito en el capítulo 10: en la web, '
    'el avatar y el vestidor 3D siguen siendo la experiencia principal; la superposición '
    'con cámara Android es experimental y se limita a vestidos. La aplicación no afirma '
    'que las mallas reproduzcan el ajuste exacto de un producto de marca.'
)
doc.add_paragraph(
    'El surtido inicial comprende 12 referencias reales identificadas en catálogos '
    'oficiales: siete vestidos y cinco faldas, distribuidos entre Nike, adidas, Levi’s '
    'y PUMA. Cada referencia tiene fotografía, variantes S/M/L y una silueta GLB '
    'referencial de su tipo. En total hay 36 variantes y 36 asociaciones de modelo. '
    'Las tres camisetas de prueba anteriores se retiraron del catálogo público sin '
    'borrar su historial. El detalle, la fuente y el archivo de imagen de cada '
    'referencia figuran en catalog-assets/women/catalog.json.'
)

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for cell, value in zip(table.rows[0].cells, ['Marca', 'Vestidos', 'Faldas']):
    cell.text = value
for brand, dresses, skirts in [('Nike', 2, 1), ('adidas', 2, 1), ('Levi’s', 1, 2), ('PUMA', 2, 1)]:
    for cell, value in zip(table.add_row().cells, [brand, str(dresses), str(skirts)]):
        cell.text = value

doc.add_paragraph(
    'Los precios en bolivianos, las existencias y los pagos son simulaciones académicas; '
    'no representan ofertas ni inventario de los fabricantes. Las fotografías provienen '
    'de sus páginas oficiales y requieren revisar permisos para uso comercial. '
    'La semilla idempotente scripts/seed-women.ts conserva cuentas, pedidos e historial. '
    'La verificación del despliegue confirmó 12 imágenes accesibles, 36 variantes con '
    'modelo, filtros de cinco faldas y siete vestidos, y respuesta correcta de los seis '
    'archivos GLB de vestido/falda por talla. Pasaron 75 pruebas de API, seis de geometría '
    'móvil y cinco del worker 3D.'
)

doc.save(PATH)
print(PATH)
